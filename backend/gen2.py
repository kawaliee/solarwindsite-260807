import django, os
os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings'); django.setup()
from apps.windsite import engine
from apps.windsite.report import build_report
from docx import Document

for label,(lat,lng,sido,sgg) in {
    '삼척 궁촌리': (37.27159528083442,129.23477162376162,'강원특별자치도','삼척시'),
    '영양 석보':   (36.6620,129.1580,'경상북도','영양군'),
}.items():
    res = engine.evaluate(lat=lat,lng=lng,radius_m=100,address=label,
                          capacity_mw=20,sido=sido,sigungu=sgg)
    raw = next((i.raw for i in res.analysis_items if '전력계통' in i.item_name), {})
    subs = raw.get('substations') or []
    print('=== %s — 조회 변전소 %d개소' % (label, len(subs)))
    for s in subs[:10]:
        print('   %-22s %5skV  %8.2fkm' % (s['name'][:22],
              (s.get('voltage') or 0)//1000 or '-', s['distance_m']/1000))
    blob = build_report(res, sido=sido, sigungu=sgg, with_maps=False)
    d = Document(__import__('io').BytesIO(blob))
    for tb in d.tables:
        hdr=[c.text.strip() for c in tb.rows[0].cells]
        if hdr and hdr[0]=='변전소':
            print('  → 보고서 6절 표 %d행' % (len(tb.rows)-1))
            for row in tb.rows[1:]:
                print('     ', [c.text.strip() for c in row.cells][:3])
    show=False
    for p in d.paragraphs:
        t=p.text.strip()
        if t.startswith('6.'): show=True; continue
        if show and t.startswith('7.'): break
        if show and t: print('  문구:', t[:200])
    print()
