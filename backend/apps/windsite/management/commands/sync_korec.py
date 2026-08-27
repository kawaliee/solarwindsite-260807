"""
전기위원회 자료 수집 → 엑셀 정리

    python manage.py sync_korec                  # 전부 (회의록 OCR 포함)
    python manage.py sync_korec --no-minutes     # 개최결과·허가대장만 (빠름)
    python manage.py sync_korec --rounds 20      # 최근 20회차 회의록만 OCR
    python manage.py sync_korec --excel-only     # 받아 둔 PDF로 엑셀만 다시

내려받은 PDF는 `data/korec/pdf/`에 그대로 둔다. 다시 돌릴 때 이미 있는 것은
건너뛰므로 회차가 늘어도 새 것만 받는다. 원본을 남기는 이유는 **판독 결과가
의심스러울 때 사람이 직접 열어 볼 수 있어야 하기 때문**이다.

⚠️ 회의록은 본문이 이미지라 OCR을 돌린다. 반쪽당 약 1.1초 × 논리 28쪽이므로
   회차당 30초쯤 걸린다. 전 회차를 처음 받으면 30분을 넘길 수 있다.
"""
from __future__ import annotations

import re
import time
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand
from django.utils import timezone

from apps.windsite import korec

DATA = Path(settings.BASE_DIR) / 'data' / 'korec'
PDF = DATA / 'pdf'
XLSX = DATA / '전기위원회_사례.xlsx'
#: 보고서가 읽는 간추린 자료. 엑셀은 사람이 보는 것이고, 이쪽은 기계가 읽는다.
#: 매번 800쪽 PDF를 다시 뜯을 수는 없기 때문이다.
CASES = DATA / 'cases.json'

#: 공지사항에서 골라낼 자료. 파일명으로 가른다.
_MINUTES = re.compile(r'회의록')
_REGISTER = re.compile(r'3MW\s*초과.*허가대장')
#: 허가취소 공고. **처분**(취소 확정)과 **청문**(취소하겠다는 예고)이 섞여 있어
#: 판독기가 둘을 갈라 준다 — 섞으면 취소 건수가 부풀려진다.
_CANCEL = re.compile(r'허가\s*취소')

#: 의결을 네 갈래로 묶어 센다. 유형 이름은 얼개마다 달라도 뜻은 이 넷이다.
GROUP_ORDER = ['가결', '조건부', '보류', '부결']
SOURCE_LABEL = {'SOLAR': '태양광', 'WIND': '풍력', 'FUELCELL': '연료전지',
                'BESS': 'BESS', 'BIO': '바이오', 'HYDRO': '수력',
                'THERMAL': '화력·복합', 'OTHER': '기타'}
KIND_LABEL = {'NEW': '신규', 'CHANGE': '변경', 'TRANSFER': '양수', 'EXTEND': '연장',
              'CANCEL': '취소', 'OTHER': '기타'}


class Command(BaseCommand):
    help = '전기위원회 개최결과·회의록·3MW초과 허가대장을 받아 엑셀로 정리한다'

    def add_arguments(self, p):
        p.add_argument('--no-results', action='store_true', help='개최결과 건너뜀')
        p.add_argument('--no-minutes', action='store_true', help='회의록 건너뜀(빠름)')
        p.add_argument('--no-register', action='store_true', help='허가대장 건너뜀')
        p.add_argument('--no-cancel', action='store_true', help='허가취소 공고 건너뜀')
        p.add_argument('--no-ocr', action='store_true',
                       help='회의록 OCR 없이 텍스트 층만 (본문 대부분이 빠진다)')
        p.add_argument('--rounds', type=int, default=0,
                       help='회의록을 최근 N회차만 판독 (0=전부)')
        p.add_argument('--excel-only', action='store_true',
                       help='새로 받지 않고 있는 PDF로 엑셀만 다시 만든다')

    def handle(self, *a, **o):
        PDF.mkdir(parents=True, exist_ok=True)
        t0 = time.time()

        results = [] if o['no_results'] else self._results(o['excel_only'])
        register = [] if o['no_register'] else self._register(o['excel_only'])
        cancels = [] if o['no_cancel'] else self._cancels(o['excel_only'])
        minutes = [] if o['no_minutes'] else self._minutes(
            o['excel_only'], not o['no_ocr'], o['rounds'])

        if results and minutes:
            hit = korec.join_minutes(results, minutes)
            self.stdout.write(self.style.MIGRATE_HEADING(
                f'\n■ 회차 매칭  {hit:,}/{len(results):,}안건에 회의록을 붙였습니다'))

        self._excel(results, minutes, register, cancels)
        self._cases(results, register, cancels)
        self.stdout.write(self.style.SUCCESS(
            f'\n완료 {time.time() - t0:.0f}초 · {XLSX}'))

    # ── 개최결과 ──────────────────────────────────────────────────────
    def _results(self, excel_only: bool) -> list[dict]:
        self.stdout.write(self.style.MIGRATE_HEADING('\n■ 위원회 개최결과'))
        files = [] if excel_only else [
            f for f in korec.crawl_files(korec.RESULT_BOARD) if f['round']]
        if files:
            self.stdout.write(f'  게시물 {len(files)}건')
        rows: list[dict] = []
        for f in files:
            self._grab(f)
        empty: list[int] = []
        for path in sorted(PDF.glob('결과_제*')):
            rd = int(korec.ROUND_RE.search(path.name).group(1))
            try:
                text, got = korec.read_result(path.read_bytes(), rd)
            except Exception as exc:                       # noqa: BLE001
                self.stderr.write(f'  ! {path.name}: {exc}')
                continue
            if not got:
                # 발전사업 안건이 없는 회차도 있다(한전 약관 변경인가만 다룬
                # 회차 등). 판독 실패와 구별해 적어 둔다 — 조용히 비면
                # 자료가 없는 것인지 못 읽은 것인지 알 수 없다.
                empty.append(rd)
            for g in got:
                g['posted'] = _posted(path.name)
            rows += got
        self.stdout.write(
            f'  판독 {len(rows):,}안건 / {len(set(r["round"] for r in rows))}회차')
        if empty:
            self.stdout.write(self.style.WARNING(
                f'  안건 0건 {len(empty)}회차: {empty}'))
        return rows

    # ── 회의록 ────────────────────────────────────────────────────────
    def _minutes(self, excel_only: bool, ocr: bool, limit: int) -> list[dict]:
        self.stdout.write(self.style.MIGRATE_HEADING('\n■ 위원회 회의록'))
        if not excel_only:
            files = [f for f in korec.crawl_files(korec.NOTICE_BOARD)
                     if f['round'] and _MINUTES.search(f['file_name'])]
            self.stdout.write(f'  게시물 {len(files)}건')
            for f in sorted(files, key=lambda x: -x['round']):
                self._grab(f, prefix='회의록')

        paths = sorted(PDF.glob('회의록_제*'),
                       key=lambda p: -int(korec.ROUND_RE.search(p.name).group(1)))
        if limit:
            paths = paths[:limit]
        # OCR은 PDF 회차에만 든다. HWP 회의록은 글자가 그대로 들어 있어 빠르다.
        pdfs = sum(1 for p in paths if p.read_bytes()[:4] == b'%PDF')
        if ocr and pdfs:
            self.stdout.write(self.style.WARNING(
                f'  PDF {pdfs}건은 본문이 이미지라 OCR합니다 — 회차당 30초쯤'))

        rows: list[dict] = []
        for path in paths:
            rd = int(korec.ROUND_RE.search(path.name).group(1))
            t = time.time()
            try:
                blob = path.read_bytes()
                pages = (korec.read_minutes(blob) if ocr
                         else [korec.read_document(blob)])
                got = korec.parse_minutes(pages, rd)
            except Exception as exc:                       # noqa: BLE001
                self.stderr.write(f'  ! 제{rd}차: {exc}')
                continue
            rows += got
            self.stdout.write(f'  제{rd}차 {len(pages)}쪽 → {len(got)}안건 '
                              f'({time.time() - t:.0f}s)')
        self.stdout.write(f'  판독 {len(rows)}안건')
        return rows

    # ── 허가대장 ──────────────────────────────────────────────────────
    def _register(self, excel_only: bool) -> list[dict]:
        self.stdout.write(self.style.MIGRATE_HEADING('\n■ 3MW 초과 발전사업 허가대장'))
        if not excel_only:
            for f in korec.crawl_files(korec.NOTICE_BOARD):
                if _REGISTER.search(f['file_name']):
                    self._grab(f, prefix='허가대장', stamp=True)
        paths = sorted(PDF.glob('허가대장_*.pdf'))
        if not paths:
            self.stdout.write(self.style.WARNING('  받은 대장이 없습니다'))
            return []
        latest = paths[-1]                    # 분기마다 갱신된다 — 최신만 쓴다
        self.stdout.write(f'  {latest.name}')
        rows = korec.parse_register(latest.read_bytes())
        first = [r for r in rows if not r['changed']]
        self.stdout.write(
            f'  판독 {len(rows):,}행 (최초허가 {len(first):,} · 변경 {len(rows) - len(first):,})')
        return rows

    # ── 허가취소 공고 ─────────────────────────────────────────────────
    def _cancels(self, excel_only: bool) -> list[dict]:
        """
        허가대장에는 **허가된 것만** 남는다. 그래서 대장만 보면 "허가만 받으면
        된다"는 그림이 된다. 실제로는 준비기간 안에 착공하지 못해 취소되는
        사업이 공고 한 건에 수십 건씩 실린다. 그것을 여기서 모은다.
        """
        self.stdout.write(self.style.MIGRATE_HEADING('\n■ 허가취소 공고'))
        if not excel_only:
            files = [f for f in korec.crawl_files(korec.NOTICE_BOARD)
                     if _CANCEL.search(f['file_name'])]
            self.stdout.write(f'  게시물 {len(files)}건')
            for f in files:
                self._grab(f, prefix='허가취소', stamp=True,
                           tag=_slug(f['file_name']))

        rows: list[dict] = []
        blank: list[str] = []
        for path in sorted(PDF.glob('허가취소_*')):
            try:
                text = korec.read_document(path.read_bytes())
                got = korec.parse_cancellation(text, path.stem)
            except Exception as exc:                       # noqa: BLE001
                self.stderr.write(f'  ! {path.name[:48]}: {exc}')
                blank.append(path.name)
                continue
            if not got:
                blank.append(path.name)
            for g in got:
                g['posted'] = _posted_stamp(path.name)
                g['file'] = path.name
            rows += got
        disp = [r for r in rows if r['kind'] == 'DISPOSAL']
        self.stdout.write(
            f'  판독 {len(rows)}건 (취소 처분 {len(disp)} · 청문 예고 '
            f'{len(rows) - len(disp)})')
        # 공고 서식이 해마다 달라 못 읽는 것이 남는다. **몇 건이 빠졌는지
        # 밝힌다** — 조용히 빠지면 "취소된 사업이 없다"로 읽힌다.
        if blank:
            self.stdout.write(self.style.WARNING(
                f'  서식이 달라 못 읽은 공고 {len(blank)}건 — 직접 열어 보십시오:'))
            for b in blank:
                self.stdout.write(f'      {b}')
        return rows

    # ── 내려받기 ──────────────────────────────────────────────────────
    def _grab(self, f: dict, prefix: str = '결과', stamp: bool = False,
              tag: str = ''):
        # 게시일을 파일명에 남긴다. 다시 돌릴 때 받은 것을 건너뛰면서도
        # 언제 것인지 알 수 있어야 하기 때문이다. 허가취소 공고는 회차가 없어
        # 같은 날 여러 건이 올라오면 겹치므로 제목 조각(tag)을 덧붙인다.
        stem = (f'{prefix}_{f["posted"]}{tag}' if stamp
                else f'{prefix}_제{f["round"]:04d}차_{f["posted"]}')
        # 이미 받아 둔 것은 확장자가 무엇이든 건너뛴다
        if any(PDF.glob(stem + '.*')):
            return
        try:
            blob = korec.download(f['file_path'], f['file_name'])
            # 게시판이 PDF·HWP·HWPX를 섞어 올린다. 이름이 아니라 **내용**으로
            # 확장자를 정해야 나중에 열린다.
            name = stem + korec.suffix_of(blob)
            dest = PDF / name
            dest.write_bytes(blob)
            self.stdout.write(f'  + {name} ({dest.stat().st_size:,}B)')
        except Exception as exc:                           # noqa: BLE001
            self.stderr.write(f'  ! {f["file_name"][:44]}: {exc}')

    # ── 조회용 자료 ───────────────────────────────────────────────────
    def _cases(self, results, register, cancels):
        import json

        # 부분 수집으로 덮어써 자료를 지우는 일을 막는다. `--no-register`로
        # 돌린 뒤 보고서에서 허가 사례가 통째로 사라지면 원인을 찾기 어렵다.
        if not (results or register or cancels):
            self.stdout.write(self.style.WARNING(
                '  판독된 것이 없어 cases.json은 그대로 둡니다'))
            return
        old = {}
        if CASES.exists():
            try:
                old = json.loads(CASES.read_text(encoding='utf-8'))
            except Exception:                              # noqa: BLE001
                old = {}
        data = _cases_json(results, register, cancels, _today())
        for key in ('register', 'agenda', 'cancel'):
            if not data[key] and old.get(key):
                data[key] = old[key]           # 이번에 안 받은 갈래는 지키다
        CASES.write_text(json.dumps(data, ensure_ascii=False),
                         encoding='utf-8')
        self.stdout.write(
            f'  cases.json — 허가 {len(data["register"]):,} · '
            f'안건 {len(data["agenda"]):,} · 취소 {len(data["cancel"]):,} '
            f'({CASES.stat().st_size / 1024:,.0f} KB)')

    # ── 엑셀 ──────────────────────────────────────────────────────────
    def _excel(self, results: list[dict], minutes: list[dict],
               register: list[dict], cancels: list[dict] | None = None):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        wb.remove(wb.active)
        hdr_fill = PatternFill('solid', fgColor='1F3864')
        hdr_font = Font(color='FFFFFF', bold=True)

        def sheet(title, cols, rows, widths):
            ws = wb.create_sheet(title)
            ws.append(cols)
            for c in ws[1]:
                c.fill, c.font = hdr_fill, hdr_font
                c.alignment = Alignment(horizontal='center', vertical='center')
            for r in rows:
                ws.append(r)
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            return ws

        # ① 허가대장 — 사례 대조의 몸통
        sheet('허가대장(3MW초과)',
              ['허가연도', '번호', '사업자', '발전소 위치', '시·도', '시·군·구',
               '에너지원', '원동력(원문)', '용량(MW)', '허가·변경일', '사업준비기간',
               '변경사항', '구분'],
              [[r['year'], r['no'], r['company'], r['location'], r['sido'],
                r['sigungu'], SOURCE_LABEL.get(r['source'], r['source']),
                r['source_raw'], r['capacity_mw'], r['permit_date'],
                r['ready_until'], r['changed'],
                '최초허가' if not r['changed'] else '변경']
               for r in register],
              [9, 6, 24, 40, 10, 12, 10, 16, 10, 13, 13, 30, 9])

        # ② 회차별 안건 — 부결·보류·사유가 실린다
        res_sorted = sorted(results, key=lambda r: (-(r['round'] or 0), r['no']))
        ws = sheet('회차별 안건',
                   ['회차', '게시일', '의결군', '의결', '연번', '사업자', '안건명',
                    '에너지원', '구분', '지역', '사유', '회의록 심의내용'],
                   [[r['round'], r.get('posted', ''), r.get('group', ''),
                     r['verdict'], r['no'], r.get('company', ''), r['title'],
                     SOURCE_LABEL.get(r['source'], r['source']),
                     KIND_LABEL.get(r['kind'], r['kind']), r['region'],
                     r['reason'], r.get('minutes', '')]
                    for r in res_sorted],
                   [7, 12, 9, 13, 6, 22, 46, 10, 8, 10, 46, 90])
        for row in ws.iter_rows(min_row=2, min_col=11, max_col=12):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')

        # ③ 회차 요약 — 어떤 회차에 무엇이 있는지 한눈에
        rounds = sorted({r['round'] for r in results if r['round']}, reverse=True)
        mrounds = {m['round'] for m in minutes}
        summary = []
        for rd in rounds:
            got = [r for r in results if r['round'] == rd]
            summary.append([
                rd, next((g.get('posted', '') for g in got if g.get('posted')), ''),
                len(got),
                sum(1 for g in got if g['source'] == 'SOLAR'),
                sum(1 for g in got if g['source'] == 'WIND'),
                *[sum(1 for g in got if g.get('group') == v) for v in GROUP_ORDER],
                '있음' if rd in mrounds else '',
            ])
        sheet('회차 요약',
              ['회차', '게시일', '안건 수', '태양광', '풍력', *GROUP_ORDER, '회의록'],
              summary, [7, 12, 9, 8, 8, *[10] * len(GROUP_ORDER), 9])

        # ④ 회의록 발췌 — 왜 그렇게 의결됐는가
        ws2 = sheet('회의록 발췌',
                    ['회차', '의결', '연번', '안건명', '에너지원', '심의 내용'],
                    [[m['round'], m.get('verdict', ''), m['no'], m['title'],
                      SOURCE_LABEL.get(m['source'], m['source']), m['text']]
                     for m in sorted(minutes,
                                     key=lambda x: (-(x['round'] or 0), x['no']))],
                    [7, 13, 6, 46, 10, 120])
        for row in ws2.iter_rows(min_row=2, min_col=6, max_col=6):
            row[0].alignment = Alignment(wrap_text=True, vertical='top')

        # ⑤ 허가취소 — 허가를 받고도 좌초한 사업
        sheet('허가취소',
              ['구분', '처분일', '게시일', '근거 회차', '사업명', '사업자',
               '사업장소', '시·도', '시·군·구', '에너지원', '용량(MW)',
               '허가번호', '처분사유'],
              [[('취소 처분' if c['kind'] == 'DISPOSAL' else '청문 예고'),
                c['disposed_on'], c.get('posted', ''), c['round'], c['name'],
                c['company'], c['location'], c['sido'], c['sigungu'],
                SOURCE_LABEL.get(c['source'], c['source']), c['capacity_mw'],
                c['permit_no'], c['reason']]
               for c in sorted(cancels or [],
                               key=lambda x: (x.get('posted', ''), x['name']),
                               reverse=True)],
              [10, 12, 12, 10, 36, 22, 40, 8, 12, 10, 11, 11, 30])

        # ⑥ 지역별 집계 — 검토 지역에 선례가 있는지 바로 보기 위한 것
        agg: dict[tuple, dict] = {}
        for r in register:
            if not r['sigungu']:
                continue
            k = (r['sido'], r['sigungu'], SOURCE_LABEL.get(r['source'], r['source']))
            a = agg.setdefault(k, {'n': 0, 'first': 0, 'mw': 0.0, 'max': 0.0,
                                   'lo': '', 'hi': ''})
            a['n'] += 1
            if not r['changed']:
                a['first'] += 1
                a['mw'] += r['capacity_mw'] or 0
                a['max'] = max(a['max'], r['capacity_mw'] or 0)
            d = r['permit_date']
            if d:
                a['lo'] = min(a['lo'], d) if a['lo'] else d
                a['hi'] = max(a['hi'], d)
        sheet('지역별 집계',
              ['시·도', '시·군·구', '에너지원', '허가 사업 수', '전체 행(변경 포함)',
               '합계 용량(MW)', '최대 용량(MW)', '최초 허가', '최근 허가'],
              [[*k, v['first'], v['n'], round(v['mw'], 1), v['max'], v['lo'], v['hi']]
               for k, v in sorted(agg.items(), key=lambda x: -x[1]['first'])],
              [10, 12, 10, 13, 18, 14, 14, 12, 12])

        DATA.mkdir(parents=True, exist_ok=True)
        wb.save(XLSX)
        self.stdout.write(
            f'\n  엑셀 6시트 — 허가대장 {len(register):,} · 안건 {len(results):,} · '
            f'회차 {len(rounds)} · 회의록 {len(minutes):,} · '
            f'허가취소 {len(cancels or []):,} · 지역 {len(agg):,}')


def _cases_json(results, register, cancels, stamp: str) -> dict:
    """
    보고서가 읽을 만큼만 간추린다.

    엑셀에는 사람이 읽을 것을 다 담지만, 조회용 자료는 **판정에 쓰는 열만**
    담는다. 회의록 전문이나 변경 이력까지 넣으면 파일이 수십 MB가 되어 매
    검토마다 읽는 값이 아니게 된다.

    허가대장은 **최초허가만** 담는다. 변경 행까지 담으면 한 사업이 여러 번
    세어져 "이 지역에 사례가 많다"는 잘못된 그림이 된다.
    """
    return {
        'built_at': stamp,
        'register': [
            {'name': r['company'], 'company': r['company'],
             'location': r['location'], 'sido': r['sido'],
             'sigungu': r['sigungu'], 'source': r['source'],
             'capacity_mw': r['capacity_mw'], 'permit_date': r['permit_date'],
             'ready_until': r['ready_until']}
            for r in register if not r['changed']],
        'agenda': [
            {'round': a['round'], 'posted': a.get('posted', ''),
             'group': a.get('group', ''), 'verdict': a['verdict'],
             'title': a['title'], 'company': a.get('company', ''),
             'source': a['source'], 'kind': a['kind'],
             'reason': a.get('reason', ''),
             # 회의록은 앞머리만 — 사업주체·위치·용량이 여기 실린다
             'minutes': (a.get('minutes') or '')[:600]}
            for a in results],
        'cancel': [
            {'kind': c['kind'], 'name': c['name'], 'company': c['company'],
             'location': c['location'], 'sido': c['sido'],
             'sigungu': c['sigungu'], 'source': c['source'],
             'capacity_mw': c['capacity_mw'], 'permit_no': c['permit_no'],
             'reason': c['reason'], 'disposed_on': c['disposed_on']}
            for c in cancels],
    }


def _slug(file_name: str) -> str:
    """공고 제목 앞머리를 파일명 꼬리에 붙인다. 같은 날 여러 건이 올라온다."""
    s = re.sub(r'_w\d+\.\w+$', '', file_name)
    s = re.sub(r'[^가-힣A-Za-z0-9]+', '', s)[:14]
    return f'_{s}' if s else ''


def _today() -> str:
    return timezone.localdate().isoformat()


def _posted_stamp(name: str) -> str:
    """허가취소 파일명에서 게시일을 되꺼낸다 — `허가취소_2026-03-18_밤실산….hwpx`"""
    m = re.search(r'_((?:19|20)\d\d-\d\d-\d\d)', name)
    return m.group(1) if m else ''


def _posted(name: str) -> str:
    """파일명 꼬리에 박아 둔 게시일을 다시 꺼낸다."""
    m = re.search(r'_((?:19|20)\d\d-\d\d-\d\d)\.pdf$', name)
    return m.group(1) if m else ''
