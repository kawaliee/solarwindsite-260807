"""
풍력 입지타당성 검토 보고서(docx) 생성기
---------------------------------------------------------------
검토 결과(EvaluationResult)를 실무 보고서 형식의 Word 문서로 만든다.
지도 4종은 검토에 실제 사용한 벡터 데이터를 재조회해 렌더링한다.

설계 원칙 (검토 엔진과 동일)
  - 추측 금지: 미확인 항목은 '확인 필요'로 그대로 싣고, 점수로 덮지 않는다.
  - 근거 동반: 모든 판정에 법령·조문·출처·검증수준을 함께 표기한다.
  - 한계 명시: 자동 판정의 한계와 다음 조치를 별도 장으로 남긴다.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime

from . import geo, maps
from .providers.vworld import VworldClient
from .schemas import EvaluationResult

logger = logging.getLogger(__name__)

STATUS_LABEL = {
    'POSSIBLE': '가능', 'CONDITIONAL': '조건부 가능',
    'IMPOSSIBLE': '불가', 'UNKNOWN': '확인 필요',
}
CONFIDENCE_LABEL = {
    'HIGH': '원문 확인', 'MEDIUM': '교차 확인', 'LOW': '미검증',
}
DIFFICULTY_LABEL = {
    'LOW': '낮음', 'MEDIUM': '보통', 'HIGH': '높음', 'CRITICAL': '매우 높음',
}


# ======================================================================
def build_report(result: EvaluationResult, *, sido: str = '', sigungu: str = '',
                 with_maps: bool = True) -> bytes:
    """검토 결과 → docx 바이트"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    _set_base_style(doc, Pt)

    site = result.site_info
    lat, lng = site.coordinates.lat, site.coordinates.lng

    # ── 표지 ────────────────────────────────────────────────────────
    t = doc.add_heading('풍력발전 입지타당성 검토 보고서', level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(site.address or f'{lat:.5f}, {lng:.5f}')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f'작성일 {datetime.now().strftime("%Y-%m-%d")}')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        '※ 본 보고서는 공공 공간정보를 자동 조회해 작성한 1차 검토 자료입니다. '
        '법적 판단이 아니며, 인허가 가능 여부는 소관기관 협의와 현장 조사로 확정해야 합니다.'
    )

    # ── 1. 검토 개요 ────────────────────────────────────────────────
    doc.add_heading('1. 검토 개요', level=1)
    _kv_table(doc, [
        ('대상지', site.address or '-'),
        ('좌표 (WGS84)', f'위도 {lat:.6f} / 경도 {lng:.6f}'),
        ('행정구역', f'{sido} {sigungu}'.strip() or '미확인'),
        ('검토 반경', f'{site.radius_m:,} m'),
        ('검토 반경 면적', f'{site.total_area_m2:,.0f} ㎡'),
        ('검토 일시', result.evaluated_at),
    ])

    # ── 2. 종합 판정 ────────────────────────────────────────────────
    doc.add_heading('2. 종합 판정', level=1)
    of = result.overall_feasibility
    _kv_table(doc, [
        ('종합 등급', STATUS_LABEL.get(of.grade, of.grade)),
        ('참고 점수', f'{of.score} / 100'),
        ('요약', of.summary),
    ])
    doc.add_paragraph(
        '점수는 항목별 상태·난이도를 감점 방식으로 합산한 **상대적 리스크 지표**입니다. '
        '미확인(UNKNOWN) 항목도 감점 대상이므로, 점수가 낮다고 해서 곧바로 부적합을 뜻하지 않습니다.'
    )

    counts: dict[str, int] = {}
    for i in result.analysis_items:
        counts[i.status.value] = counts.get(i.status.value, 0) + 1
    _table(doc, ['판정', '항목 수'],
           [[STATUS_LABEL.get(k, k), str(v)] for k, v in counts.items()])

    # ── 3. 필지·가용면적 ────────────────────────────────────────────
    doc.add_heading('3. 필지 및 가용면적', level=1)
    parcel = _raw_of(result, '필지·지적 분석')
    if parcel.get('by_jimok'):
        _kv_table(doc, [
            ('조회 필지 수', f'{parcel.get("parcel_count", 0):,} 필지'),
            ('지적 면적 합계', f'{parcel.get("total_parcel_area_m2", 0):,.0f} ㎡'),
            ('가용 지목 면적 (임야·잡종지·목장용지)',
             f'{parcel.get("usable_area_m2", 0):,.0f} ㎡'),
            ('공공용지 성격 면적 (도로·하천 등)',
             f'{parcel.get("excluded_area_m2", 0):,.0f} ㎡'),
            ('전용 절차 필요 면적',
             ', '.join(f'{k} {v:,.0f}㎡' for k, v in
                       (parcel.get('conversion_needed') or {}).items()) or '없음'),
        ])
        _table(doc, ['지목', '면적(㎡)', '필지 수'],
               [[k, f'{v["area_m2"]:,.0f}', f'{v["count"]:,}']
                for k, v in parcel['by_jimok'].items()])
        if parcel.get('truncated'):
            doc.add_paragraph('⚠️ 조회 상한에 도달해 일부 필지가 누락되었을 수 있습니다.')
    else:
        doc.add_paragraph('필지 정보를 조회하지 못했습니다. 상세 사유는 9장을 참고하십시오.')

    if with_maps:
        _insert_map(doc, Cm, _cadastral_map(lat, lng, site.radius_m),
                    '[지도 1] 지적 현황도')

    # ── 4. 규제 저촉 현황 ───────────────────────────────────────────
    doc.add_heading('4. 규제 저촉 현황', level=1)
    regs = [i for i in result.analysis_items
            if i.status.value in ('IMPOSSIBLE', 'CONDITIONAL')]
    if regs:
        _table(doc, ['항목', '판정', '난이도', '근거 법령·조문', '검증'],
               [[i.item_name, STATUS_LABEL.get(i.status.value, ''),
                 DIFFICULTY_LABEL.get(i.difficulty.value, ''),
                 f'{i.law} {i.article}'.strip() or '-',
                 CONFIDENCE_LABEL.get(i.confidence.value, '')] for i in regs])
        for i in regs:
            doc.add_heading(f'4.{regs.index(i) + 1} {i.item_name}', level=2)
            doc.add_paragraph(i.reason)
            if i.action_required:
                doc.add_paragraph(f'▷ 필요 조치: {i.action_required}')
    else:
        doc.add_paragraph('저촉이 확인된 규제 항목이 없습니다.')

    if with_maps:
        _insert_map(doc, Cm, _regulation_map(result, lat, lng, site.radius_m),
                    '[지도 2] 규제 구역 중첩도')

    # ── 5. 환경·산림·재해 ───────────────────────────────────────────
    doc.add_heading('5. 환경·산림·재해', level=1)
    env = [i for i in result.analysis_items if i.category in ('환경', '산림')]
    if env:
        _table(doc, ['항목', '판정', '내용'],
               [[i.item_name, STATUS_LABEL.get(i.status.value, ''), i.reason[:400]]
                for i in env])
    else:
        doc.add_paragraph('해당 항목이 없습니다.')

    # ── 6. 전력계통 연계 ────────────────────────────────────────────
    doc.add_heading('6. 전력계통 연계', level=1)
    grid = _raw_of(result, '전력계통 연계(변전소·송전선로)')
    subs = grid.get('substations') or []
    if subs:
        _table(doc, ['변전소', '전압(kV)', '직선거리', '출처'],
               [[s['name'], str((s.get('voltage') or 0) // 1000 or '-'),
                 geo.format_distance(s['distance_m']), s.get('osm', '') or '수기']
                for s in subs[:8]])
        doc.add_paragraph(
            '⚠️ 위 목록은 OpenStreetMap 기반 위치 탐색 결과이며 **접속 가능 용량(계통 여유도)은 '
            '포함되지 않습니다.** 한전 계통연계 사전검토를 반드시 별도로 신청해야 합니다.'
        )
    else:
        doc.add_paragraph('변전소가 조회되지 않았습니다. 한전ON에서 직접 확인이 필요합니다.')

    if with_maps:
        _insert_map(doc, Cm, _surroundings_map(lat, lng, site.radius_m),
                    '[지도 3] 주변 현황도 (건물·도로)')

    # ── 7. 이격거리 동심원 분석 ─────────────────────────────────────
    doc.add_heading('7. 이격거리 동심원 분석', level=1)
    quiet = _raw_of(result, '정온시설 이격거리(동심원 분석)')
    rings = quiet.get('rings') or []
    facs = quiet.get('facilities') or []
    if rings:
        _table(doc, ['이격 대상', '조례 기준(m)', '기준 내 시설 수', '근거 조례'],
               [[r['target'], f'{r["distance_m"]:,}', str(r['count']),
                 r.get('ordinance', '')] for r in rings])
    else:
        doc.add_paragraph('해당 지자체의 이격거리 조례가 DB에 등록되어 있지 않아 '
                          '저촉 여부를 판정하지 않았습니다.')
    if facs:
        _table(doc, ['정온시설', '유형', '직선거리'],
               [[f['name'], f['kind'], geo.format_distance(f['distance_m'])]
                for f in facs[:10]])
        doc.add_paragraph(
            '※ 본 분석은 **대상 지점 1점 기준**입니다. 실제 이격거리는 개별 발전기 위치를 '
            '확정한 뒤 재산출해야 하며, OSM 미등재 주거지가 있을 수 있어 현장 실사가 필요합니다.'
        )

    if with_maps and (rings or facs):
        _insert_map(doc, Cm, _setback_map(lat, lng, rings, facs),
                    f'[지도 4] 이격거리 동심원도 (근접 {len(facs)}개소 표시)')

    # ── 8. 인허가 로드맵 ────────────────────────────────────────────
    doc.add_heading('8. 인허가 로드맵', level=1)
    steps = [s for s in result.permit_roadmap if s.applicable]
    if steps:
        _table(doc, ['순서', '단계', '절차명', '소관기관', '근거', '법정기간(일)'],
               [[str(s.order), s.phase, s.name, s.authority,
                 f'{s.law} {s.article}'.strip(),
                 str(s.statutory_days) if s.statutory_days else '-'] for s in steps])
    else:
        doc.add_paragraph('인허가 절차 데이터가 없습니다. `seed_windsite`를 실행하십시오.')

    # ── 9. 확인 필요 사항 및 한계 ───────────────────────────────────
    doc.add_heading('9. 확인 필요 사항 및 자동 판정의 한계', level=1)
    if result.data_gaps:
        for g in result.data_gaps:
            doc.add_paragraph(g, style='List Bullet')
    else:
        doc.add_paragraph('추가 확인이 필요한 항목이 없습니다.')
    doc.add_paragraph(
        '본 검토는 공개 공간정보의 조회 결과에 기반합니다. 다음 항목은 공개 API가 없어 '
        '자동 판정되지 않으며 기관 협의가 필요합니다 — 군사기지·비행안전구역 및 레이더 전파영향, '
        '전력계통 접속 가능 용량, 허브고도 풍황.'
    )

    # ── 10. 근거 법령 ───────────────────────────────────────────────
    doc.add_heading('10. 근거 법령', level=1)
    laws = result.applicable_laws or []
    if laws:
        _table(doc, ['법령', '분류', '역할', '검증'],
               [[l.get('name', ''), l.get('category', ''),
                 (l.get('purpose', '') or '')[:80],
                 CONFIDENCE_LABEL.get(l.get('confidence', ''), '')] for l in laws])
        doc.add_paragraph(
            '검증란이 "미검증"인 항목은 법령 원문 대조가 완료되지 않은 상태입니다. '
            '실무 적용 전 국가법령정보센터에서 원문을 확인하십시오.'
        )

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ======================================================================
# 문서 헬퍼
# ======================================================================
def _set_base_style(doc, Pt) -> None:
    style = doc.styles['Normal']
    style.font.name = '맑은 고딕'
    style.font.size = Pt(10)


def _table(doc, headers: list[str], rows: list[list[str]]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row[:len(headers)]):
            cells[i].text = str(v)
    doc.add_paragraph()
    return t


def _kv_table(doc, pairs: list[tuple[str, str]]):
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Light List Accent 1'
    for k, v in pairs:
        cells = t.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v)
    doc.add_paragraph()
    return t


def _insert_map(doc, Cm, png: bytes | None, caption: str) -> None:
    if not png:
        doc.add_paragraph(f'{caption} — 렌더링에 필요한 데이터를 조회하지 못했습니다.')
        return
    doc.add_picture(io.BytesIO(png), width=Cm(16))
    p = doc.add_paragraph(caption)
    p.runs[0].italic = True


def _raw_of(result: EvaluationResult, item_name: str) -> dict:
    for i in result.analysis_items:
        if i.item_name == item_name:
            return i.raw or {}
    return {}


# ======================================================================
# 지도 데이터 재조회
# ======================================================================
def _fetch_geoms(layer_id: str, lat: float, lng: float, radius_m: int,
                 name_field: str = '') -> list[dict]:
    """레이어에서 도형을 받아 계량좌표계 shapely 도형 목록으로 돌려준다."""
    try:
        feats, _ = VworldClient.fetch_all(layer_id, lat, lng, radius_m, max_pages=3)
    except Exception:                                           # noqa: BLE001
        logger.exception('지도용 %s 조회 실패', layer_id)
        return []

    out: list[dict] = []
    for f in feats:
        g = geo.geom_from_geojson(f.get('geometry'))
        if g is None:
            continue
        try:
            gm = geo.to_metric(g)
        except Exception:                                       # noqa: BLE001
            continue
        props = f.get('properties') or {}
        out.append({'geom': gm, 'name': (props.get(name_field) or '') if name_field else '',
                    'props': props})
    return out


def _layer_id(code: str, fallback: str) -> str:
    from .models import RegulationLayer
    l = RegulationLayer.objects.filter(code=code, is_active=True).first()
    return l.layer_id if l else fallback


def _cadastral_map(lat: float, lng: float, radius_m: int) -> bytes | None:
    from .providers.cadastral import parse_jimok

    if not geo.GEO_AVAILABLE:
        return None
    rows = _fetch_geoms(_layer_id('연속지적', 'lp_pa_cbnd_bubun'), lat, lng, radius_m)
    if not rows:
        return None
    feats = [{'geom': r['geom'], 'jimok': parse_jimok(r['props'].get('jibun', '')) or '미상'}
             for r in rows]
    try:
        return maps.cadastral_map(lat, lng, radius_m, feats)
    except Exception:                                           # noqa: BLE001
        logger.exception('지적 현황도 렌더링 실패')
        return None


def _regulation_map(result: EvaluationResult, lat: float, lng: float,
                    radius_m: int) -> bytes | None:
    from .models import RegulationLayer

    if not geo.GEO_AVAILABLE:
        return None
    hit_codes = [(i.raw or {}).get('code') for i in result.analysis_items
                 if i.status.value in ('IMPOSSIBLE', 'CONDITIONAL') and (i.raw or {}).get('code')]
    if not hit_codes:
        return None

    layers: list[dict] = []
    for l in RegulationLayer.objects.filter(code__in=hit_codes, is_active=True):
        rows = _fetch_geoms(l.layer_id, lat, lng,
                            radius_m + l.search_margin_m + l.proximity_m)
        if not rows:
            continue
        site = geo.point_metric(lat, lng)
        layers.append({
            'name': l.title or l.code,
            'geoms': [r['geom'] for r in rows],
            'overlapping': any(site.distance(r['geom']) == 0 for r in rows),
        })
    if not layers:
        return None
    try:
        return maps.regulation_map(lat, lng, radius_m, layers)
    except Exception:                                           # noqa: BLE001
        logger.exception('규제 중첩도 렌더링 실패')
        return None


def _surroundings_map(lat: float, lng: float, radius_m: int) -> bytes | None:
    if not geo.GEO_AVAILABLE:
        return None
    r = max(radius_m * 2, 1000)
    buildings = [x['geom'] for x in
                 _fetch_geoms(_layer_id('건물', 'lt_c_spbd'), lat, lng, r)]
    roads = [x['geom'] for x in
             _fetch_geoms(_layer_id('도로', 'lt_l_sprd'), lat, lng, r)]
    if not buildings and not roads:
        return None
    try:
        return maps.surroundings_map(lat, lng, radius_m, buildings, roads)
    except Exception:                                           # noqa: BLE001
        logger.exception('주변 현황도 렌더링 실패')
        return None


def _setback_map(lat: float, lng: float, rings: list[dict],
                 facilities: list[dict]) -> bytes | None:
    if not geo.GEO_AVAILABLE or not (rings or facilities):
        return None
    try:
        return maps.setback_map(lat, lng, rings or [{'target': '참고', 'distance_m': 1000}],
                                facilities)
    except Exception:                                           # noqa: BLE001
        logger.exception('이격거리 동심원도 렌더링 실패')
        return None
