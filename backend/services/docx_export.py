"""
Word(docx) 산출물 생성 — 계약서 초안/검토 결과 Word 파일 내보내기
"""
import os
import uuid
import logging
from pathlib import Path
from django.conf import settings

logger = logging.getLogger(__name__)


def export_contract_draft(draft) -> str:
    """
    계약서 초안 → Word 파일
    Returns: 저장된 파일 경로
    """
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    doc = Document()

    # 스타일 설정
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Malgun Gothic'
    font.size = Pt(11)

    # 제목
    title_para = doc.add_heading(draft.title or draft.template.name_ko, level=0)
    title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # 부제
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = subtitle.add_run(f'({draft.template.name_en or ""})')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph('')

    # 본문 — generated_content를 마크다운 형식에서 변환
    content = draft.generated_content or ''
    lines = content.split('\n')

    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph('')
            continue

        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith('*') and line.endswith('*'):
            para = doc.add_paragraph()
            run = para.add_run(line.strip('*'))
            run.italic = True
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        else:
            doc.add_paragraph(line)

    # 하단 안내
    doc.add_paragraph('')
    footer = doc.add_paragraph()
    run = footer.add_run('※ 본 계약서는 재생E AI Agent에 의해 자동 생성된 초안이며, 법률 검토가 필요합니다.')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.italic = True

    # 저장
    output_dir = Path(settings.MEDIA_ROOT) / 'exports'
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = f'draft_{uuid.uuid4().hex[:8]}.docx'
    filepath = output_dir / filename
    doc.save(str(filepath))

    logger.info(f'Contract draft exported: {filepath}')
    return str(filepath)


def export_contract_review(review) -> str:
    """
    계약서 검토 결과 → Word 파일
    Returns: 저장된 파일 경로
    """
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor, Cm
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    doc = Document()

    # 스타일 설정
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Malgun Gothic'
    font.size = Pt(11)

    # 제목
    title_para = doc.add_heading(f'계약서 검토 보고서', level=0)
    title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # 정보
    doc.add_paragraph(f'검토 대상: {review.title}')
    if review.template:
        doc.add_paragraph(f'계약 유형: {review.template.name_ko}')
    doc.add_paragraph(f'검토 지시: {review.review_instruction}')
    doc.add_paragraph('')

    # 총평
    if review.summary:
        doc.add_heading('검토 총평', level=1)
        doc.add_paragraph(review.summary)
        doc.add_paragraph('')

    # 검토 결과 표
    doc.add_heading('조항별 검토 결과', level=1)

    findings = review.findings.all().order_by('order_index')

    if findings.exists():
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Light Grid Accent 1'

        # 헤더
        headers = ['조항', '위험도', '지적 내용', '수정 방향']
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        # 데이터 행
        severity_map = {'high': '독소', 'mid': '불리', 'low': '경고/누락'}
        for finding in findings:
            row = table.add_row()
            row.cells[0].text = finding.clause_ref or '—'
            row.cells[1].text = severity_map.get(finding.severity, finding.severity)
            row.cells[2].text = finding.finding
            row.cells[3].text = finding.suggestion or ''

    # 하단 안내
    doc.add_paragraph('')
    footer = doc.add_paragraph()
    run = footer.add_run('※ 본 검토 보고서는 재생E AI Agent에 의해 자동 생성되었으며, 법률 전문가 검토가 필요합니다.')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.italic = True

    # 저장
    output_dir = Path(settings.MEDIA_ROOT) / 'exports'
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = f'review_{uuid.uuid4().hex[:8]}.docx'
    filepath = output_dir / filename
    doc.save(str(filepath))

    logger.info(f'Contract review exported: {filepath}')
    return str(filepath)
